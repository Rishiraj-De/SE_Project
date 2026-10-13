import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import datetime

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_srs_document(filename="Digital_Archaeologist_SRS_Document.docx"):
    doc = Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("SOFTWARE REQUIREMENTS SPECIFICATION (IEEE 830) | DIGITAL ARCHAEOLOGIST")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(100, 116, 139)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Department of Computer Science & Engineering • Academic Project Submission • Page ")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(100, 116, 139)

    # Base styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # ------------------ COVER PAGE ------------------
    cover_box = doc.add_table(rows=1, cols=1)
    cover_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_cell = cover_box.cell(0, 0)
    set_cell_background(c_cell, "0F172A")
    set_cell_margins(c_cell, top=360, bottom=360, left=360, right=360)
    
    cp = c_cell.paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = cp.add_run("SOFTWARE REQUIREMENTS SPECIFICATION (SRS)\nACCORDING TO IEEE STD 830-1998\n\n")
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(148, 163, 184)

    run_title = cp.add_run("DIGITAL ARCHAEOLOGIST\n")
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(255, 255, 255)

    run_desc = cp.add_run("An Intelligent Static Analysis & Software Maintenance Platform for Legacy Codebase Deconstruction\n")
    run_desc.font.size = Pt(13)
    run_desc.font.italic = True
    run_desc.font.color.rgb = RGBColor(203, 213, 225)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    # Metadata Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Document Title", "Software Requirements Specification (SRS) - Digital Archaeologist"),
        ("Project Version", "Version 1.0.0 (Production Release)"),
        ("Standard Compliance", "IEEE Recommended Practice for SRS (IEEE Std 830-1998) / ISO/IEC/IEEE 29148"),
        ("Prepared By", "Student Project Team / Engineering Department"),
        ("Department & University", "Department of Computer Science and Engineering"),
        ("Submission Date", datetime.datetime.now().strftime("%B %d, %Y"))
    ]

    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "FFFFFF")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)
        
        p0 = cell_lbl.paragraphs[0]
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(15, 23, 42)

        p1 = cell_val.paragraphs[0]
        r1 = p1.add_run(val)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_page_break()

    # Helpers
    def add_h1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)
        return h

    def add_h2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)
        return h

    def add_h3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(11)
        r.font.bold = True
        r.font.color.rgb = RGBColor(71, 85, 105)
        return h

    def add_callout(text, title="NOTE / SPECIFICATION DETAIL:"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        p = cell.paragraphs[0]
        r_title = p.add_run(f"{title}\n")
        r_title.font.bold = True
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = RGBColor(37, 99, 235)
        r_text = p.add_run(text)
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = RGBColor(51, 65, 85)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ------------------ TABLE OF CONTENTS ------------------
    add_h1("TABLE OF CONTENTS")
    toc_items = [
        ("1. INTRODUCTION", "3"),
        ("   1.1 Purpose & Scope", "3"),
        ("   1.2 Document Conventions", "3"),
        ("   1.3 Intended Audience & Reading Suggestions", "4"),
        ("   1.4 Project Scope & Objectives", "4"),
        ("   1.5 References & Industry Standards", "4"),
        ("2. OVERALL DESCRIPTION", "5"),
        ("   2.1 Product Perspective", "5"),
        ("   2.2 Product Functions (High-Level)", "5"),
        ("   2.3 User Classes and Characteristics", "6"),
        ("   2.4 Operating Environment", "6"),
        ("   2.5 Design & Implementation Constraints", "6"),
        ("   2.6 Assumptions and Dependencies", "7"),
        ("3. SYSTEM FEATURES & FUNCTIONAL REQUIREMENTS", "8"),
        ("   3.1 FR-01: Multi-Stage Ingestion & Real-Time Pipeline Tracker", "8"),
        ("   3.2 FR-02: Multi-Language AST Parsing & Symbol Extraction", "9"),
        ("   3.3 FR-03: McCabe Cyclomatic Complexity Computation", "9"),
        ("   3.4 FR-04: Dependency Graph Construction & DFS Cycle Detection", "10"),
        ("   3.5 FR-05: Git History Mining & Volatile File Churn Dynamics", "11"),
        ("   3.6 FR-06: Static Blast Radius & Transitive Impact Analyzer", "12"),
        ("   3.7 FR-07: Technical Debt Catalog & Remediation Advisory", "13"),
        ("   3.8 FR-08: Risk Analysis: Churn vs. Complexity Scatter Matrix", "14"),
        ("   3.9 FR-09: Contextual AI Assistant with Grounded Citations", "15"),
        ("   3.10 FR-10: Automated Executive Documentation & Audit Reports", "16"),
        ("   3.11 FR-11: VS-Code Inspired Monaco Codebase Explorer", "17"),
        ("   3.12 FR-12: Interactive React Flow Architecture Canvas", "17"),
        ("4. EXTERNAL INTERFACE REQUIREMENTS", "18"),
        ("   4.1 User Interfaces (UI/UX Design Rules)", "18"),
        ("   4.2 Hardware Interfaces", "18"),
        ("   4.3 Software Interfaces (API Contract)", "19"),
        ("   4.4 Communication Interfaces", "19"),
        ("5. NON-FUNCTIONAL REQUIREMENTS (NFRs)", "20"),
        ("   5.1 Performance Requirements", "20"),
        ("   5.2 Safety & Security Requirements", "20"),
        ("   5.3 Software Quality Attributes (Reliability, Availability, Maintainability, Usability)", "21"),
        ("6. DATA DESIGN & SYSTEM ARCHITECTURE", "22"),
        ("   6.1 Entity-Relationship (ER) Model & Relational Tables", "22"),
        ("   6.2 Data Flow Diagrams (Level 0, Level 1)", "23"),
        ("7. VERIFICATION & ACCEPTANCE CRITERIA", "24"),
    ]

    for item, pg in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(item)
        r1.font.size = Pt(9.5)
        if item.startswith("1.") or item.startswith("2.") or item.startswith("3.") or item.startswith("4.") or item.startswith("5.") or item.startswith("6.") or item.startswith("7."):
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(15, 23, 42)
        else:
            r1.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_page_break()

    # ------------------ SECTION 1: INTRODUCTION ------------------
    add_h1("1. INTRODUCTION")
    
    add_h2("1.1 Purpose")
    doc.add_paragraph(
        "This Software Requirements Specification (SRS) document details the complete functional and non-functional "
        "specifications for 'Digital Archaeologist', an intelligent legacy-code analysis and software maintenance platform. "
        "The purpose of this document is to establish a rigorous, formal specification adhering to the IEEE Std 830-1998 "
        "guidelines for academic evaluation, peer review, and industrial engineering implementation."
    )

    add_h2("1.2 Document Conventions")
    doc.add_paragraph(
        "This document follows standard IEEE notation conventions. Requirements are uniquely identified using prefixes: "
        "'FR-xx' designates a Functional Requirement, while 'NFR-xx' designates a Non-Functional Requirement. "
        "Keywords indicating specification level follow RFC 2119: SHALL indicates mandatory compliance, "
        "SHOULD indicates recommended practice, and MAY indicates optional extensibility."
    )

    add_h2("1.3 Intended Audience & Reading Suggestions")
    doc.add_paragraph(
        "The primary audience includes Academic Project Evaluators, Professors, Senior Software Engineers, Systems Architects, "
        "and Quality Assurance engineers. Evaluators interested in system requirements and architectural soundness "
        "should focus on Section 3 (System Features) and Section 6 (System Architecture). Engineers implementing external "
        "integrations should review Section 4 (Interface Requirements)."
    )

    add_h2("1.4 Project Scope & Objectives")
    doc.add_paragraph(
        "Legacy software maintenance accounts for an estimated 60% to 80% of total lifetime software engineering costs. "
        "When engineering teams inherit undocumented or aging repositories, they face steep cognitive hurdles: undocumented "
        "circular dependencies, hidden blast radiuses during modification, high-churn volatility, and dead code accumulation. "
        "Existing commercial tools either produce overwhelming static code metric dumps without interactive context, or rely on "
        "generative AI chatbots that hallucinate non-existent files and lines. "
        "Digital Archaeologist addresses this critical gap by providing a deterministic, AST-driven forensic environment. "
        "The platform ingests source repositories, computes McCabe Cyclomatic Complexity, constructs directed dependency graphs, "
        "isolates architectural cycles via Depth-First Search (DFS), calculates transitive Blast Radius impact, and presents "
        "grounded explanations with verified file citations and line numbers."
    )

    add_h2("1.5 References & Industry Standards")
    refs = [
        "IEEE Std 830-1998: IEEE Recommended Practice for Software Requirements Specifications.",
        "McCabe, T. J. (1976). 'A Complexity Measure'. IEEE Transactions on Software Engineering, SE-2(4), 308-320.",
        "Tarjan, R. (1972). 'Depth-First Search and Linear Graph Algorithms'. SIAM Journal on Computing, 1(2), 146-160.",
        "Fowler, M. (2018). 'Refactoring: Improving the Design of Existing Code' (2nd ed.). Addison-Wesley Professional.",
        "ISO/IEC/IEEE 29148:2018: Systems and software engineering — Life cycle processes — Requirements engineering."
    ]
    for r in refs:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(r)

    # ------------------ SECTION 2: OVERALL DESCRIPTION ------------------
    add_h1("2. OVERALL DESCRIPTION")

    add_h2("2.1 Product Perspective")
    doc.add_paragraph(
        "Digital Archaeologist is a standalone, full-stack, distributed intelligence system. It does not replace code compilers "
        "or continuous integration (CI) runners; rather, it functions as an upstream Architectural Observability and Pre-Refactoring "
        "Platform. The system operates in a three-tier topology consisting of: (1) an editorial, high-density Next.js web application, "
        "(2) a high-throughput Express.js REST API service, and (3) a high-speed Python Abstract Syntax Tree (AST) & Git mining engine, "
        "backed by a normalized SQLite relational database managed via Prisma ORM."
    )

    add_h2("2.2 Product Functions (High-Level Overview)")
    doc.add_paragraph("The core capabilities delivered by Digital Archaeologist encompass:")
    funcs = [
        "Repository Ingestion: Real-time onboarding of Git URLs and pre-bundled legacy repositories with live multi-stage telemetry.",
        "Static AST Parsing: Lexical analysis extracting classes, methods, functions, and import statements without code execution.",
        "Cyclomatic Complexity Scoring: McCabe decision point evaluation across all control flow branches.",
        "Dependency Graph Construction: Directed adjacency graph modeling module coupling, fan-in (dependents), and fan-out (depends-on).",
        "DFS Cycle Detection: Three-state coloring graph algorithm isolating reciprocal circular dependencies (e.g. ServiceA <-> ServiceB).",
        "Git Forensics & Churn Dynamics: Longitudinal mining of commit history, author signatures, and file volatility frequencies.",
        "Risk Quadrant Mapping: Interactive scatter plotting correlating Change Frequency (Churn) against Cyclomatic Complexity.",
        "Static Blast Radius Analyzer: Transitive call tree derivation identifying downstream modules, database writes, and API side-effects.",
        "Technical Debt Workspace: Multi-category debt inventory with severity badges and concrete remediation strategies.",
        "Contextual AI Assistant: Deterministic, citation-backed conversational interface citing exact files, line numbers, and commit hashes.",
        "Executive Audit Dossier Generator: Automated Markdown documentation and printable executive compliance assessment reports."
    ]
    for f in funcs:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        p.add_run(f)

    add_h2("2.3 User Classes and Characteristics")
    doc.add_paragraph(
        "1. Lead Software Architects: Require global system topology visualization, circular dependency isolation, and technical debt prioritization.\n"
        "2. Senior Software Engineers: Require blast radius impact calculation before modifying critical functions, and Git churn hotspot awareness.\n"
        "3. Junior Engineers & Interns: Require codebase exploration, contextual AI explanations with exact line-level citations, and architectural documentation.\n"
        "4. Academic Evaluators & Professors: Assess methodological rigor, McCabe complexity calculations, and architectural modularity."
    )

    add_h2("2.4 Operating Environment")
    doc.add_paragraph(
        "• Host OS: Cross-platform support across Microsoft Windows 10/11 (x64), Linux (Ubuntu 20.04+, Debian, Fedora), and macOS (Apple Silicon / Intel).\n"
        "• Web Browsers: Modern evergreen browsers supporting ECMAScript 2022+ and WebAssembly (Google Chrome 110+, Microsoft Edge 110+, Mozilla Firefox 115+, Apple Safari 16+).\n"
        "• Execution Runtimes: Node.js v18.0.0 or higher, Python 3.9 or higher, and Git CLI 2.30+."
    )

    add_h2("2.5 Design & Implementation Constraints")
    doc.add_paragraph(
        "• Zero Untrusted Code Execution: Security constraint mandating that source files are analyzed statically via AST tokenization without executing binaries or invoking dynamic code eval.\n"
        "• Offline Capability: System must function in isolated network environments using pre-warmed AST caches and local deterministic heuristic models.\n"
        "• Single-Page High Density Layout: VS Code and Linear inspired editorial user interface with high information density, avoiding unnecessary decorative animations."
    )

    # ------------------ SECTION 3: FUNCTIONAL REQUIREMENTS ------------------
    add_h1("3. SYSTEM FEATURES & FUNCTIONAL REQUIREMENTS")

    frs = [
        ("FR-01", "Repository Ingestion & Real-Time Pipeline Tracker",
         "The system SHALL accept a public Git repository URL, authenticated private repository token, or local directory path. "
         "Upon ingestion, the system SHALL execute and display a 9-stage pipeline progress indicator: (1) Repository Cloned, "
         "(2) Languages Detected, (3) Source Tree Indexed, (4) AST Generated, (5) Dependencies Extracted, (6) Git History Analyzed, "
         "(7) Architecture Reconstructed, (8) Risk Analyzed, and (9) Documentation Generated.",
         "Git URL string or bundled demo selection.",
         "Background clone execution, progress telemetry emission, database entity initialization.",
         "Persisted Repository and AnalysisRun entities, complete pipeline completion signal."),

        ("FR-02", "Multi-Language AST Parsing & Symbol Extraction",
         "The analysis engine SHALL recursively parse source files matching supported extensions (.ts, .tsx, .js, .jsx, .py, .java, .go, .rs). "
         "For each file, the parser SHALL extract: Total Lines of Code (LOC), Class Definitions, Function/Method Definitions with exact starting line numbers, "
         "and Raw Import / Require statements.",
         "Source code files on filesystem.",
         "Lexical tokenization and regular expression pattern recognition ignoring comment blocks and string literals.",
         "Structured JSON array of module definitions, symbol catalogs, and language distribution counts."),

        ("FR-03", "McCabe Cyclomatic Complexity Computation",
         "The system SHALL compute McCabe Cyclomatic Complexity V(G) for each module and function according to the formula: "
         "V(G) = 1 + P, where P represents the count of predicate decision points. The parser SHALL count instances of: "
         "'if', 'else if', 'elif', 'for', 'while', 'case', 'catch', 'except', ternary operators ('? :'), logical AND ('&&'), and logical OR ('||').",
         "Source file text content.",
         "Multi-pattern branch matching across control flow constructs.",
         "Integer complexity value assigned to each module; modules with complexity > 15 SHALL be flagged as Technical Debt."),

        ("FR-04", "Dependency Graph Construction & DFS Cycle Detection",
         "The system SHALL resolve relative and package import statements into a directed graph structure comprising both forward adjacency "
         "(u -> v, indicating u depends on v) and reverse adjacency (v -> u, indicating u is a dependent of v). "
         "The engine SHALL execute a Depth-First Search (DFS) algorithm with a three-color recursion stack (0=unvisited, 1=visiting, 2=visited). "
         "If a node with state 1 is encountered during traversal, the cycle SHALL be isolated and recorded with full module sequence chains.",
         "Module collection and raw import lists.",
         "Path normalization, directory hierarchy resolution, recursive DFS cycle detection.",
         "Adjacency graph data structure and array of circular dependencies flagged with 'HIGH' severity."),

        ("FR-05", "Git History Mining & Volatile File Churn Dynamics",
         "The engine SHALL execute Git CLI subcommands ('git log --pretty=format:%h|%an|%ad|%s' and 'git log --name-only') to extract historical commit telemetry. "
         "The system SHALL aggregate commit frequency per file to compute 'File Churn'. The engine SHALL sort files descending by change frequency "
         "to identify the most volatile files across repository lifetime.",
         "Repository .git directory.",
         "Log parsing, author extraction, commit hash identification, churn count aggregation.",
         "Chronological commit history list, unique contributor roster, and ranked Most Volatile Files table."),

        ("FR-06", "Static Blast Radius & Transitive Impact Analyzer",
         "Given a target symbol or module name (e.g. 'PaymentService.processPayment'), the system SHALL recursively traverse the reverse adjacency "
         "graph to identify all direct and transitive callers. The engine SHALL inspect call paths for database write operations (SQL writes) "
         "and third-party external gateway interactions (e.g., Stripe, PayPal, banking APIs). The system SHALL compute an overall Risk Rating "
         "(CRITICAL, HIGH, MEDIUM, LOW) and render an interactive call hierarchy tree.",
         "Target symbol string and target file path.",
         "Recursive graph traversal, side-effect pattern scanning, heuristic risk weighting.",
         "BlastRadiusResult payload including affected module count, database write operations, external endpoints, and tree hierarchy."),

        ("FR-07", "Technical Debt Catalog & Remediation Advisory",
         "The system SHALL automatically synthesize detected code smells into a categorized, actionable Technical Debt catalog. "
         "Categories SHALL include: High Complexity (>15), Circular Dependencies, Dead Code Candidates (unreferenced utility functions), "
         "and Insecure / Deprecated APIs (e.g. MD5 hashing in authentication paths). Each item SHALL include severity rating, component name, "
         "file path, exact issue description, and specific engineering remediation recommendations.",
         "Analysis results (modules, cycles, complexity).",
         "Rule-based classification and recommendation synthesis.",
         "Structured tabular Technical Debt catalog with filter and export capabilities."),

        ("FR-08", "Risk Analysis: Churn vs. Complexity Scatter Matrix",
         "The system SHALL synthesize an interactive two-dimensional risk matrix mapping Historical Change Frequency (X-axis) against "
         "Cyclomatic Complexity (Y-axis). Modules situated in the upper-right quadrant (High Churn AND High Complexity) SHALL be classified "
         "into the 'Critical Hotspot' category and highlighted with visual warning badges.",
         "Module complexity data and Git file churn counts.",
         "Coordinate calculation, quadrant boundary evaluation, risk score synthesis ((Complexity * 2) + (Churn * 3)).",
         "Interactive scatter plot visualization with interactive point inspection drawers."),

        ("FR-09", "Contextual AI Assistant with Grounded Citations",
         "The system SHALL provide an interactive contextual software maintenance assistant. For any user architectural query, "
         "the assistant SHALL provide an authoritative technical answer accompanied by verified citations. "
         "Each citation SHALL contain: exact relative file path, 1-indexed line number, verbatim source code snippet, and related Git commit hash. "
         "The assistant SHALL prohibit ungrounded hallucinations.",
         "User natural language prompt and active repository index context.",
         "Deterministic AST-grounded retrieval matching symbol graph with heuristic response templates.",
         "AIResponse payload with answer narrative, structured citation cards, and recommended follow-up questions."),

        ("FR-10", "Automated Documentation & Audit Report Generation",
         "The system SHALL generate comprehensive Markdown documentation and audit reports on demand. "
         "The Executive Assessment Report SHALL synthesize health scores, circular loop warnings, blast radius hotspots, and debt summaries. "
         "The UI SHALL provide instant one-click clipboard copy and clean PDF-printable layout capabilities.",
         "Active repository analysis data model.",
         "Dynamic Markdown document compilation with structured tables, executive scorecards, and formatted code blocks.",
         "Rendered Markdown view, copy confirmation toast, and print-stylesheet optimized layout."),

        ("FR-11", "VS-Code Inspired Monaco Codebase Explorer",
         "The system SHALL embed a Microsoft Monaco Editor instance inside a 3-pane layout: (Left) recursive directory tree, "
         "(Center) Monaco code editor with syntax highlighting, line numbers, and minimap, (Right) AST symbol inspector displaying "
         "class/method metrics with one-click Blast Radius triggers.",
         "File selection events from tree browser.",
         "Asynchronous file content retrieval from backend API, Monaco editor model configuration.",
         "Interactive code viewer with active line tracking, symbol breakdown, and direct refactoring triggers."),

        ("FR-12", "Interactive React Flow Architecture Canvas",
         "The system SHALL render the module dependency topology on an interactive React Flow canvas. "
         "Nodes SHALL be styled by architectural role (Services, Controllers, Repositories, Database Singletons). "
         "Circular dependencies SHALL be rendered with animated directional edges. "
         "Clicking any node SHALL open an inspector drawer displaying LOC, complexity, fan-in/fan-out gauges, and change frequency.",
         "Adjacency graph, module metadata, circular dependency chains.",
         "Graph layout computation, node/edge generation, role filtering.",
         "Interactive pan/zoom canvas, search filter, and detailed node inspector sidebar.")
    ]

    for code, title, desc, inp, proc, outp in frs:
        add_h2(f"{code}: {title}")
        doc.add_paragraph(desc)
        
        fr_tbl = doc.add_table(rows=3, cols=2)
        fr_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        fr_tbl.autofit = False
        
        specs = [
            ("Input Parameters", inp),
            ("Processing Logic", proc),
            ("Output Artifacts", outp)
        ]
        for s_idx, (s_lbl, s_val) in enumerate(specs):
            row = fr_tbl.rows[s_idx]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Inches(1.8)
            c1.width = Inches(4.7)
            set_cell_background(c0, "F1F5F9")
            set_cell_background(c1, "FFFFFF")
            set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
            set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
            
            p0 = c0.paragraphs[0]
            r0 = p0.add_run(s_lbl)
            r0.font.bold = True
            r0.font.size = Pt(8.5)
            r0.font.color.rgb = RGBColor(15, 23, 42)

            p1 = c1.paragraphs[0]
            r1 = p1.add_run(s_val)
            r1.font.size = Pt(8.5)
            r1.font.color.rgb = RGBColor(51, 65, 85)
            
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ------------------ SECTION 4: INTERFACE REQUIREMENTS ------------------
    add_h1("4. EXTERNAL INTERFACE REQUIREMENTS")

    add_h2("4.1 User Interfaces")
    doc.add_paragraph(
        "The web interface follows an editorial developer design system inspired by Linear and VS Code. "
        "The interface incorporates: (1) Collapsible navigation sidebar with Ctrl+B shortcut support, (2) Top telemetry header "
        "displaying current repository identifier, default branch, and overall Health Index pill, (3) Dark Slate (#0F172A) and Graphite (#1E293B) "
        "color palette with high information density, and (4) Strict semantic badge hierarchy (CRITICAL=Red, HIGH=Orange, MEDIUM=Amber, LOW=Blue)."
    )

    add_h2("4.2 Hardware Interfaces")
    doc.add_paragraph(
        "The platform does not require specialized hardware peripherals. Minimum client hardware specifications: "
        "Dual-core x86_64 or ARM64 processor, 4 GB RAM, 1080p display resolution (1920x1080 recommended for 3-pane Monaco viewer). "
        "Minimum server host specifications: 2 vCPUs, 2 GB RAM, 5 GB available disk space for temporary repository checkouts."
    )

    add_h2("4.3 Software Interfaces & API Specifications")
    doc.add_paragraph("The Express REST API server exposes the following endpoints under the /api namespace:")
    
    api_table = doc.add_table(rows=12, cols=3)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    api_headers = ["HTTP Method & Route", "Input Payload / Query", "Response Payload"]
    for idx, h in enumerate(api_headers):
        c = api_table.rows[0].cells[idx]
        set_cell_background(c, "0F172A")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    api_rows = [
        ("POST /api/analyze", "{ repoUrl: string, repoName?: string }", "AnalysisResult object with pipeline run ID"),
        ("GET /api/analysis/status", "None", "{ isAnalyzing: boolean, analysis: AnalysisResult }"),
        ("GET /api/repos/:id/overview", "repoId in path", "Summary metrics, health scores, alert list"),
        ("GET /api/repos/:id/architecture", "repoId in path", "React Flow nodes, edges, circular chains"),
        ("GET /api/repos/:id/dependencies", "repoId in path", "Coupling table with fan-in, fan-out, loops"),
        ("GET /api/repos/:id/files", "repoId in path", "File hierarchy list with LOC and symbol counts"),
        ("GET /api/repos/:id/file-content", "?path=relative_path", "{ content: string, language: string, loc: number }"),
        ("GET /api/repos/:id/git-history", "repoId in path", "Commit records, authors, most changed files"),
        ("GET /api/repos/:id/tech-debt", "repoId in path", "TechnicalDebt array with severity and remediation"),
        ("GET /api/repos/:id/risk-analysis", "repoId in path", "RiskPoint array mapping Churn vs Complexity"),
        ("POST /api/repos/:id/impact-analysis", "{ targetSymbol: string, targetFile?: string }", "BlastRadiusResult with call hierarchy tree")
    ]

    for r_idx, (m, inp, outp) in enumerate(api_rows, start=1):
        row = api_table.rows[r_idx]
        for c_idx, val in enumerate([m, inp, outp]):
            cell = row.cells[c_idx]
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_h2("4.4 Communication Interfaces")
    doc.add_paragraph(
        "Client-server communication utilizes HTTP/1.1 and HTTP/2 protocols over TCP/IP with JSON payload encoding. "
        "CORS headers are configured to permit distributed cross-origin API invocation. Standard HTTPS (TLS 1.3) "
        "is enforced in production deployments."
    )

    # ------------------ SECTION 5: NON-FUNCTIONAL REQUIREMENTS ------------------
    add_h1("5. NON-FUNCTIONAL REQUIREMENTS (NFRs)")

    nfrs = [
        ("NFR-01", "Performance & Latency",
         "The Python analysis engine SHALL complete parsing, complexity calculation, and cycle detection for repositories "
         "up to 10,000 Lines of Code (LOC) within 500 milliseconds on standard dual-core hardware. "
         "All backend REST API query endpoints SHALL return responses within 100 milliseconds for pre-warmed analysis runs. "
         "Frontend Next.js views SHALL achieve First Contentful Paint (FCP) in under 1.2 seconds."),

        ("NFR-02", "Security & Code Isolation",
         "The analysis engine SHALL strictly analyze source code through static lexical parsing without invoking runtime execution, "
         "eval(), or arbitrary process spawning of analyzed files. Ingested repository URLs SHALL be validated and sanitized "
         "to prevent shell injection vulnerabilities. Database interactions SHALL utilize parameterized Prisma queries to preclude SQL injection."),

        ("NFR-03", "Reliability & Fault Tolerance",
         "If the Python engine encounters unparseable syntax in an individual file, the engine SHALL skip the malformed file, "
         "log an error, and continue parsing remaining repository files without crashing. "
         "If the backend server is unreachable, the frontend client SHALL fall back to pre-cached demo telemetry to preserve usability."),

        ("NFR-04", "Maintainability & Modularity",
         "The codebase SHALL adhere to strict separation of concerns across presentation, orchestration, analysis, and data access layers. "
         "All TypeScript code SHALL be compiled with strict mode enabled ('strict': true). "
         "Cyclomatic complexity of individual platform methods SHALL remain under 12."),

        ("NFR-05", "Portability & Platform Independence",
         "The backend and frontend services SHALL be containerizable via standard Dockerfiles. "
         "Path handling logic SHALL utilize platform-agnostic path resolvers (Pathlib in Python, path.resolve in Node.js) "
         "to guarantee seamless operation across Windows, Linux, and macOS filesystems.")
    ]

    for code, title, desc in nfrs:
        add_h2(f"{code}: {title}")
        doc.add_paragraph(desc)

    # ------------------ SECTION 6: DATA DESIGN ------------------
    add_h1("6. DATA DESIGN & SYSTEM ARCHITECTURE")

    add_h2("6.1 Relational Database Architecture (Prisma / SQLite)")
    doc.add_paragraph(
        "The persistence layer is governed by a relational database schema containing 15 normalized entities. "
        "Primary models include:"
    )

    models_data = [
        ("Repository", "Represents the target codebase (id, name, url, defaultBranch, ownerId)."),
        ("AnalysisRun", "Represents an individual analysis execution (id, branch, commitHash, healthScore, startedAt, completedAt)."),
        ("File", "Stores file-level metrics (path, name, language, loc, complexity, analysisRunId)."),
        ("ClassDefinition", "Tracks identified classes and line numbers within each file."),
        ("FunctionDefinition", "Tracks function/method definitions, parameter signatures, and line numbers."),
        ("Dependency", "Models directed dependency links between source and destination files with isCircular flag."),
        ("CircularDependency", "Records isolated circular dependency cycles with complete loop chain descriptions."),
        ("TechnicalDebt", "Stores actionable debt items with severity ratings, issues, categories, and remediation advice."),
        ("RiskScore", "Stores calculated risk scores synthesizing cyclomatic complexity with Git change frequency."),
        ("Commit & CommitChange", "Maintains chronological Git commit records, author metadata, and file modification counts.")
    ]

    for m_name, m_desc in models_data:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{m_name}: ")
        r.font.bold = True
        p.add_run(m_desc)

    add_h2("6.2 Data Flow Architecture")
    doc.add_paragraph(
        "Level 0 Context Diagram:\n"
        "[Software Engineer / User] ---> (Git Repo URL / Selection) ---> [Digital Archaeologist Platform] ---> (Interactive Visualizations & Dossiers)\n\n"
        "Level 1 Pipeline Subsystem Flow:\n"
        "1. Ingestion Controller receives Git target and clones repository to local workspace.\n"
        "2. Python Analysis Subprocess parses source code, computes McCabe complexity, builds graph, and mines Git logs.\n"
        "3. Results are serialized to JSON and persisted via Prisma ORM into SQLite relational database.\n"
        "4. Express API Controller exposes domain endpoints for overview, architecture, dependencies, git churn, and risk.\n"
        "5. Next.js Client queries REST endpoints, populating React Flow canvas, Monaco Editor, and Recharts scatter plot."
    )

    # ------------------ SECTION 7: ACCEPTANCE CRITERIA ------------------
    add_h1("7. VERIFICATION & ACCEPTANCE CRITERIA")

    doc.add_paragraph(
        "The software platform shall be deemed acceptable for production deployment and academic certification upon satisfying "
        "the following formal criteria:"
    )

    acc_table = doc.add_table(rows=7, cols=3)
    acc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    acc_headers = ["Verification Item", "Test Condition", "Pass Criteria"]
    for idx, h in enumerate(acc_headers):
        c = acc_table.rows[0].cells[idx]
        set_cell_background(c, "0F172A")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    acc_rows = [
        ("AST Tokenization Accuracy", "Analyze bundled legacy-payment-system codebase", "All classes, methods, and files correctly parsed with 100% symbol detection."),
        ("McCabe Complexity Accuracy", "Evaluate PaymentService.ts with nested branching", "Calculated complexity matches manual decision point count (V(G)=18)."),
        ("Cycle Detection Verification", "Evaluate PaymentService <-> TransactionService circular import", "DFS algorithm flags cycle with severity=HIGH and renders animated red edge."),
        ("Git Churn Identification", "Mine 4-year commit history of demo repository", "PaymentService.ts correctly ranked #1 in volatility (6 historical modifications)."),
        ("Blast Radius Computation", "Execute impact analysis on PaymentService.processPayment()", "Returns 3 affected modules, 1 DB write operation, and external gateway warning."),
        ("Deterministic AI Citation", "Query assistant regarding PaymentService dependency", "Returns verbatim snippet, line 82 citation, and commit hash 65a6a0e.")
    ]

    for r_idx, (item, cond, crit) in enumerate(acc_rows, start=1):
        row = acc_table.rows[r_idx]
        for c_idx, val in enumerate([item, cond, crit]):
            cell = row.cells[c_idx]
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8)
            if c_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Sign-off box
    sign_box = doc.add_table(rows=2, cols=2)
    sign_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    s0 = sign_box.cell(0, 0)
    s1 = sign_box.cell(0, 1)
    s0.width = Inches(3.2)
    s1.width = Inches(3.2)
    set_cell_margins(s0, top=100, bottom=100, left=100, right=100)
    set_cell_margins(s1, top=100, bottom=100, left=100, right=100)
    
    s0.paragraphs[0].add_run("PREPARED AND SUBMITTED BY:\n\n___________________________________\nStudent Signature / Roll Number\nDepartment of Computer Science & Engineering").font.size = Pt(8.5)
    s1.paragraphs[0].add_run("APPROVED AND ACCEPTED BY:\n\n___________________________________\nFaculty Project Guide / Professor\nDepartment of Computer Science & Engineering").font.size = Pt(8.5)

    doc.save(filename)
    print(f"Successfully generated {filename}")

if __name__ == "__main__":
    create_srs_document()
